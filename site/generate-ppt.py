#!/usr/bin/env python3
"""Generate Bilibili promotional video PPT for Second Brain plugin.

Light theme with purple accents. Introduces Obsidian first,
then plugin features. No WeChat contact.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import os

# Light theme
BG_DARK = RGBColor(0xF8, 0xF7, 0xFC)        # light lavender bg
BG_CARD = RGBColor(0xEF, 0xED, 0xF6)         # card bg
BG_CARD_ACTIVE = RGBColor(0xE8, 0xE4, 0xF3)  # active card
OBS_PURPLE = RGBColor(0x7C, 0x3A, 0xED)      # brand purple
OBS_PURPLE_BRIGHT = RGBColor(0x6D, 0x28, 0xD9) # deeper purple
OBS_PURPLE_DIM = RGBColor(0xA7, 0x8B, 0xFA)  # light purple
BRAND_PURPLE = RGBColor(0x7C, 0x3A, 0xED)    # brand purple
BRAND_GREEN = RGBColor(0x05, 0x96, 0x69)     # dark green
TEXT_WHITE = RGBColor(0x37, 0x41, 0x51)       # dark gray text (primary)
TEXT_BRIGHT = RGBColor(0x1F, 0x29, 0x37)      # near-black text (headings)
TEXT_DIM = RGBColor(0x6B, 0x72, 0x80)         # secondary text
TEXT_MUTED = RGBColor(0x9C, 0xA3, 0xAF)       # muted text
RED = RGBColor(0xDC, 0x26, 0x26)              # red
BORDER_DIM = RGBColor(0xDD, 0xDB, 0xE5)       # light border

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

W = prs.slide_width
H = prs.slide_height


def set_shape_transparency(shape, pct):
    """Set fill transparency percentage (0-100)."""
    fill = shape.fill._fill
    srgb = fill.find(qn('a:solidFill'))
    if srgb is not None:
        color_elem = srgb[0]
        alpha = color_elem.makeelement(qn('a:alpha'), {})
        alpha.set('val', str(int((100 - pct) * 1000)))
        color_elem.append(alpha)


def add_bg(slide, color=BG_DARK):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text_box(slide, left, top, width, height, text, font_size=24,
                 color=TEXT_WHITE, bold=False, alignment=PP_ALIGN.LEFT,
                 font_name="Microsoft YaHei"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_multiline_text(slide, left, top, width, height, lines, font_size=28,
                       color=TEXT_WHITE, bold=False, line_spacing=1.5,
                       alignment=PP_ALIGN.LEFT, font_name="Microsoft YaHei"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True

    for i, line_data in enumerate(lines):
        if isinstance(line_data, dict):
            text = line_data.get("text", "")
            line_color = line_data.get("color", color)
            line_bold = line_data.get("bold", bold)
            line_size = line_data.get("size", font_size)
        else:
            text = line_data
            line_color = color
            line_bold = bold
            line_size = font_size

        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()

        p.text = text
        p.font.size = Pt(line_size)
        p.font.color.rgb = line_color
        p.font.bold = line_bold
        p.font.name = font_name
        p.alignment = alignment
        p.space_after = Pt(font_size * (line_spacing - 1))

    return txBox


def add_rounded_rect(slide, left, top, width, height, fill_color=BG_CARD,
                     border_color=None, text="", font_size=18,
                     text_color=TEXT_WHITE):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()

    if text:
        tf = shape.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].text = text
        tf.paragraphs[0].font.size = Pt(font_size)
        tf.paragraphs[0].font.color.rgb = text_color
        tf.paragraphs[0].font.name = "Microsoft YaHei"
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf.paragraphs[0].space_before = Pt(0)

    return shape


def add_circle(slide, left, top, size, fill_color=OBS_PURPLE, text="",
               font_size=14, text_color=TEXT_WHITE):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, left, top, size, size
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()

    if text:
        tf = shape.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].text = text
        tf.paragraphs[0].font.size = Pt(font_size)
        tf.paragraphs[0].font.color.rgb = text_color
        tf.paragraphs[0].font.name = "Microsoft YaHei"
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf.paragraphs[0].space_before = Pt(0)

    return shape


def add_arrow(slide, left, top, width, height, color=OBS_PURPLE):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_decor_line(slide, left, top, width, color=OBS_PURPLE_DIM):
    """Add a thin horizontal decorative line."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, left, top, width, Pt(2)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


# ============================================================
# SLIDE 1: Title
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

# Decorative accent line
add_decor_line(slide, Inches(5.5), Inches(1.5), Inches(2.3), OBS_PURPLE_DIM)

add_rounded_rect(slide, Inches(5.2), Inches(1.8), Inches(2.9), Inches(0.5),
                 fill_color=RGBColor(0xDD, 0xDB, 0xE5),
                 border_color=OBS_PURPLE_DIM,
                 text="Obsidian Plugin", font_size=14, text_color=OBS_PURPLE)

add_text_box(slide, Inches(0), Inches(2.5), W, Inches(1.2),
             "Second Brain", font_size=64, color=TEXT_BRIGHT, bold=True,
             alignment=PP_ALIGN.CENTER)

add_text_box(slide, Inches(0), Inches(3.8), W, Inches(0.8),
             "AI \u9A71\u52A8\u7684\u77E5\u8BC6\u7F16\u8BD1\u5668",
             font_size=28, color=OBS_PURPLE, alignment=PP_ALIGN.CENTER)

add_text_box(slide, Inches(0), Inches(4.7), W, Inches(1.0),
             "\u628A\u6563\u4E71\u7684\u7B14\u8BB0\u4E22\u8FDB\u53BB\uFF0CAI \u81EA\u52A8\u5E2E\u4F60\u6574\u7406\u6210\u7ED3\u6784\u5316\u77E5\u8BC6\u5E93",
             font_size=20, color=TEXT_DIM, alignment=PP_ALIGN.CENTER)

add_decor_line(slide, Inches(5.5), Inches(5.8), Inches(2.3), OBS_PURPLE_DIM)

# ============================================================
# SLIDE 2: What is Obsidian?
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_text_box(slide, Inches(0), Inches(0.6), W, Inches(0.8),
             "\u5148\u804A\u804A Obsidian",
             font_size=40, color=TEXT_BRIGHT, bold=True,
             alignment=PP_ALIGN.CENTER)

add_decor_line(slide, Inches(6.0), Inches(1.4), Inches(1.3), OBS_PURPLE)

add_text_box(slide, Inches(1.5), Inches(1.8), Inches(10), Inches(0.8),
             "\u4E00\u6B3E\u514D\u8D39\u7684\u7B14\u8BB0\u8F6F\u4EF6\uFF0C\u6240\u6709\u7B14\u8BB0\u4EE5\u6587\u4EF6\u5F62\u5F0F\u5B58\u5728\u4F60\u7684\u7535\u8111\u4E0A\uFF0C\u4F60\u5B8C\u5168\u62E5\u6709\u6570\u636E",
             font_size=20, color=TEXT_DIM, alignment=PP_ALIGN.CENTER)

obs_features = [
    {"icon": "\u21C4", "title": "\u53CC\u5411\u94FE\u63A5", "desc": "\u7B14\u8BB0\u4E4B\u95F4\u53EF\u4EE5\u4E92\u76F8\u5F15\u7528\n\u50CF\u7F51\u9875\u4E00\u6837\u8DF3\u8F6C"},
    {"icon": "\u2609", "title": "\u672C\u5730\u5B58\u50A8", "desc": "\u6570\u636E\u5728\u4F60\u81EA\u5DF1\u7684\u7535\u8111\u4E0A\n\u4E0D\u4F9D\u8D56\u4E91\u670D\u52A1\uFF0C\u9690\u79C1\u5B89\u5168"},
    {"icon": "\u2726", "title": "\u63D2\u4EF6\u751F\u6001", "desc": "\u793E\u533A\u5F00\u53D1\u4E86\u5343\u4F59\u63D2\u4EF6\n\u53EF\u4EE5\u968F\u610F\u5B9A\u5236\u529F\u80FD"},
]

card_w = Inches(3.2)
card_h = Inches(3.2)
start_x = Inches(1.8)
gap = Inches(0.6)

for i, feat in enumerate(obs_features):
    x = start_x + i * (card_w + gap)
    y = Inches(3.0)
    add_rounded_rect(slide, x, y, card_w, card_h,
                     fill_color=BG_CARD, border_color=BORDER_DIM)

    # Icon circle
    add_circle(slide, x + Inches(1.15), y + Inches(0.3), Inches(0.8),
               fill_color=OBS_PURPLE_DIM, text=feat["icon"],
               font_size=24, text_color=TEXT_BRIGHT)

    add_text_box(slide, x, y + Inches(1.3), card_w, Inches(0.5),
                 feat["title"], font_size=22, color=TEXT_BRIGHT, bold=True,
                 alignment=PP_ALIGN.CENTER)

    add_text_box(slide, x + Inches(0.3), y + Inches(1.9), card_w - Inches(0.6), Inches(1.2),
                 feat["desc"], font_size=16, color=TEXT_DIM,
                 alignment=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 3: Pain points
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_text_box(slide, Inches(0), Inches(0.6), W, Inches(0.8),
             "\u4F46\u7528 Obsidian \u8D8A\u4E45\uFF0C\u4F60\u4F1A\u9047\u5230\u8FD9\u4E9B\u95EE\u9898",
             font_size=34, color=TEXT_BRIGHT, bold=True,
             alignment=PP_ALIGN.CENTER)

add_decor_line(slide, Inches(6.0), Inches(1.4), Inches(1.3), OBS_PURPLE)

pains = [
    "\u8BFB\u4E66\u7B14\u8BB0\u5199\u4E86\u51E0\u5341\u7BC7\uFF0C\u6587\u7AE0\u526A\u85CF\u649E\u4E86\u4E0A\u767E\u6761",
    "\u64AD\u5BA2\u7B14\u8BB0\u3001\u95EA\u5FF5\u788E\u7247\u6563\u843D\u5404\u5904",
    "\u5199\u7684\u65F6\u5019\u89C9\u5F97\u5B66\u5230\u4E86\uFF0C\u60F3\u7528\u7684\u65F6\u5019\u627E\u4E0D\u5230",
    "\u5168\u5C40\u641C\u7D22\u51FA\u4E00\u5806\u96F6\u6563\u7ED3\u679C\uFF0C\u8FD8\u8981\u81EA\u5DF1\u62FC\u51D1\u4E0A\u4E0B\u6587",
]

for i, pain in enumerate(pains):
    y = Inches(2.0) + i * Inches(1.05)
    add_rounded_rect(slide, Inches(2), y, Inches(9.3), Inches(0.8),
                     fill_color=BG_CARD, border_color=BORDER_DIM,
                     text=pain, font_size=20, text_color=TEXT_DIM)

# Bottom keywords
kw_y = Inches(6.2)
kw_items = ["\u627E\u4E0D\u5230", "\u8FDE\u4E0D\u4E0A", "\u7528\u4E0D\u8D77\u6765"]
kw_w = Inches(2.5)
kw_gap = Inches(0.5)
kw_start = Inches(3.2)
for i, kw in enumerate(kw_items):
    x = kw_start + i * (kw_w + kw_gap)
    add_rounded_rect(slide, x, kw_y, kw_w, Inches(0.7),
                     fill_color=RGBColor(0xFD, 0xEB, 0xEB),
                     border_color=RED,
                     text=kw, font_size=24, text_color=RED)

# ============================================================
# SLIDE 4: Existing solutions
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_text_box(slide, Inches(0), Inches(1.0), W, Inches(0.8),
             "\u73B0\u6709\u7684\u89E3\u51B3\u65B9\u6848\uFF1F",
             font_size=40, color=TEXT_BRIGHT, bold=True,
             alignment=PP_ALIGN.CENTER)

add_decor_line(slide, Inches(6.0), Inches(1.8), Inches(1.3), OBS_PURPLE)

solutions = [
    {"title": "\u624B\u52A8\u5F52\u7C7B\u6574\u7406", "result": "\u592A\u7D2F\u4E86\uFF0C\u575A\u6301\u4E0D\u4E0B\u6765"},
    {"title": "Dataview \u505A\u7D22\u5F15", "result": "\u5B66\u4E60\u6210\u672C\u592A\u9AD8"},
    {"title": "\u624B\u5199\u53CC\u94FE\u5173\u7CFB", "result": "\u6709\u6548\u4F46\u8017\u65F6\u8017\u529B"},
]

for i, sol in enumerate(solutions):
    y = Inches(2.5) + i * Inches(1.5)
    add_rounded_rect(slide, Inches(2.5), y, Inches(3.5), Inches(1.0),
                     fill_color=BG_CARD, border_color=BORDER_DIM,
                     text=sol["title"], font_size=24, text_color=TEXT_DIM)
    add_text_box(slide, Inches(6.5), y + Inches(0.15), Inches(4), Inches(0.7),
                 sol["result"], font_size=24, color=RED)

# ============================================================
# SLIDE 5: Karpathy LLM Wiki methodology
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_rounded_rect(slide, Inches(0.8), Inches(0.5), Inches(2.0), Inches(0.45),
                 fill_color=RGBColor(0xDD, 0xDB, 0xE5),
                 border_color=OBS_PURPLE_DIM,
                 text="\u8BBE\u8BA1\u7406\u5FF5", font_size=13, text_color=OBS_PURPLE)

add_text_box(slide, Inches(0.8), Inches(1.2), Inches(11), Inches(0.8),
             "\u7075\u611F\u6765\u6E90: Karpathy \u7684 LLM Wiki",
             font_size=36, color=TEXT_BRIGHT, bold=True)

add_decor_line(slide, Inches(0.8), Inches(2.0), Inches(2), OBS_PURPLE_DIM)

add_text_box(slide, Inches(0.8), Inches(2.4), Inches(11), Inches(0.8),
             "OpenAI \u524D\u7814\u7A76\u79D1\u5B66\u5BB6 Andrej Karpathy \u63D0\u51FA\u7684\u77E5\u8BC6\u7BA1\u7406\u65B9\u6CD5\uFF1A\u8BA9 AI \u6301\u7EED\u6784\u5EFA\u4E00\u4E2A\u53EF\u590D\u5229\u7684 Wiki\uFF0C\u800C\u4E0D\u662F\u6BCF\u6B21\u4ECE\u96F6\u5F00\u59CB\u68C0\u7D22",
             font_size=18, color=TEXT_DIM)

# Three-layer architecture
layers = [
    {"name": "Raw", "desc": "\u539F\u59CB\u7D20\u6750\u5C42\n\u6563\u4E71\u7684\u7B14\u8BB0\u3001\u526A\u85CF\u3001\u95EA\u5FF5"},
    {"name": "Wiki", "desc": "\u7F16\u8BD1\u8F93\u51FA\u5C42\n\u7ED3\u6784\u5316\u3001\u53CC\u94FE\u63A5\u7684\u77E5\u8BC6\u5E93"},
    {"name": "Schema", "desc": "\u5143\u6570\u636E\u5C42\n\u7D22\u5F15\u3001\u5206\u7C7B\u3001\u5065\u5EB7\u5EA6\u68C0\u67E5"},
]

layer_w = Inches(3.2)
layer_start_x = Inches(1.2)
layer_gap = Inches(0.65)
layer_arrow_w = Inches(0.6)

for i, layer in enumerate(layers):
    x = layer_start_x + i * (layer_w + layer_gap + layer_arrow_w)
    y = Inches(3.5)

    add_rounded_rect(slide, x + Inches(0.1), y, layer_w - Inches(0.2), Inches(2.2),
                     fill_color=BG_CARD_ACTIVE if i == 1 else BG_CARD,
                     border_color=OBS_PURPLE if i == 1 else BORDER_DIM)

    add_text_box(slide, x, y + Inches(0.2), layer_w, Inches(0.5),
                 layer["name"], font_size=26, color=OBS_PURPLE_BRIGHT if i == 1 else TEXT_BRIGHT,
                 bold=True, alignment=PP_ALIGN.CENTER)

    add_text_box(slide, x + Inches(0.2), y + Inches(0.9), layer_w - Inches(0.4), Inches(1.2),
                 layer["desc"], font_size=14, color=TEXT_DIM, alignment=PP_ALIGN.CENTER)

    if i < 2:
        ax = x + layer_w + Inches(0.1)
        ay = y + Inches(0.8)
        add_arrow(slide, ax, ay, layer_arrow_w - Inches(0.2), Inches(0.5), OBS_PURPLE)

add_text_box(slide, Inches(0), Inches(6.2), W, Inches(0.6),
             "Second Brain \u5C31\u662F\u8FD9\u4E2A\u7406\u5FF5\u7684 Obsidian \u843D\u5730\u5B9E\u73B0",
             font_size=20, color=OBS_PURPLE_BRIGHT, bold=True,
             alignment=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 6: One-line pitch
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_text_box(slide, Inches(0), Inches(1.0), W, Inches(0.6),
             "Second Brain",
             font_size=18, color=OBS_PURPLE, alignment=PP_ALIGN.CENTER)

add_decor_line(slide, Inches(6.0), Inches(1.6), Inches(1.3), OBS_PURPLE)

add_text_box(slide, Inches(1), Inches(2.2), Inches(11), Inches(1.2),
             "\u628A\u6563\u4E71\u7684\u7B14\u8BB0\u4E22\u8FDB\u53BB\uFF0CAI \u81EA\u52A8\u5E2E\u4F60\u6574\u7406\u6210\n\u5E26\u53CC\u5411\u94FE\u63A5\u7684\u7ED3\u6784\u5316\u77E5\u8BC6\u5E93",
             font_size=36, color=TEXT_BRIGHT, bold=True,
             alignment=PP_ALIGN.CENTER)

add_text_box(slide, Inches(1), Inches(4.0), Inches(11), Inches(1.5),
             "\u5B89\u88C5\u540E\u76F4\u63A5\u5728 Obsidian \u91CC\u4F7F\u7528\uFF0C\u4E0D\u9700\u8981\u5199\u4EE3\u7801\uFF0C\u4E0D\u9700\u8981\u6280\u672F\u80CC\u666F\n\u914D\u7F6E\u4E00\u4E2A AI API Key \u5C31\u53EF\u4EE5\u5F00\u59CB",
             font_size=20, color=TEXT_DIM, alignment=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 7: Three-step flow
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_text_box(slide, Inches(0), Inches(0.5), W, Inches(0.8),
             "\u4E09\u6B65\u6784\u5EFA\u4F60\u7684\u77E5\u8BC6\u4F53\u7CFB",
             font_size=36, color=TEXT_BRIGHT, bold=True,
             alignment=PP_ALIGN.CENTER)

add_decor_line(slide, Inches(6.0), Inches(1.3), Inches(1.3), OBS_PURPLE)

steps = [
    {"num": "01", "title": "\u653E\u5165\u7D20\u6750",
     "desc": "\u6587\u7AE0\u3001\u8BFB\u4E66\u7B14\u8BB0\u3001\u95EA\u5FF5\u7B14\u8BB0\n\u4E22\u5230 raw/ \u76EE\u5F55\u5373\u53EF"},
    {"num": "02", "title": "AI \u7F16\u8BD1",
     "desc": "AI \u81EA\u52A8\u63D0\u53D6\u6982\u5FF5\u3001\u5B9E\u4F53\u3001\u6765\u6E90\n\u751F\u6210\u5E26\u53CC\u5411\u94FE\u63A5\u7684 Wiki \u9875\u9762"},
    {"num": "03", "title": "\u6D4F\u89C8\u4E0E\u5BF9\u8BDD",
     "desc": "\u5361\u7247\u7D22\u5F15\u3001AI \u5BF9\u8BDD\u3001\u8BED\u4E49\u641C\u7D22\n\u4F60\u7684\u77E5\u8BC6\u771F\u6B63\u201C\u7528\u8D77\u6765\u201D"},
]

box_w = Inches(3.2)
start_x = Inches(1.2)
gap_x = Inches(0.65)
arrow_w = Inches(0.6)

for i, step in enumerate(steps):
    x = start_x + i * (box_w + gap_x + arrow_w)
    y = Inches(1.8)

    # Step number
    add_circle(slide, x + Inches(1.15), y, Inches(0.9),
               fill_color=OBS_PURPLE_DIM, text=step["num"],
               font_size=28, text_color=TEXT_BRIGHT)

    # Title
    add_text_box(slide, x, y + Inches(1.1), box_w, Inches(0.5),
                 step["title"], font_size=26, color=TEXT_BRIGHT, bold=True,
                 alignment=PP_ALIGN.CENTER)

    # Description card
    add_rounded_rect(slide, x + Inches(0.1), y + Inches(1.8),
                     box_w - Inches(0.2), Inches(1.8),
                     fill_color=BG_CARD_ACTIVE if i == 1 else BG_CARD,
                     border_color=OBS_PURPLE if i == 1 else BORDER_DIM,
                     text=step["desc"], font_size=16, text_color=TEXT_DIM)

    # Arrow between steps
    if i < 2:
        ax = x + box_w + Inches(0.1)
        ay = y + Inches(2.3)
        add_arrow(slide, ax, ay, arrow_w - Inches(0.2), Inches(0.5), OBS_PURPLE)

# ============================================================
# SLIDE 8: Core Feature 1 - AI Compile
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_rounded_rect(slide, Inches(0.8), Inches(0.5), Inches(2.0), Inches(0.45),
                 fill_color=RGBColor(0xDD, 0xDB, 0xE5),
                 border_color=OBS_PURPLE_DIM,
                 text="\u6838\u5FC3\u529F\u80FD", font_size=13, text_color=OBS_PURPLE)

add_text_box(slide, Inches(0.8), Inches(1.2), Inches(11), Inches(0.8),
             "AI \u7F16\u8BD1 -- \u788E\u7247\u81EA\u52A8\u53D8\u7ED3\u6784",
             font_size=36, color=TEXT_BRIGHT, bold=True)

add_decor_line(slide, Inches(0.8), Inches(2.0), Inches(2), OBS_PURPLE_DIM)

features = [
    "\u81EA\u52A8\u63D0\u53D6\u6982\u5FF5\u3001\u5B9E\u4F53\u3001\u77E5\u8BC6\u6765\u6E90",
    "\u751F\u6210\u5E26 [[\u53CC\u5411\u94FE\u63A5]] \u7684 Wiki \u9875\u9762",
    "\u81EA\u52A8\u6784\u5EFA\u5206\u7C7B\u7D22\u5F15\uFF08\u6982\u5FF5 / \u5B9E\u4F53 / \u7D20\u6750\u6458\u8981\uFF09",
    "\u589E\u91CF\u7F16\u8BD1\uFF1A\u53EA\u5904\u7406\u6709\u53D8\u5316\u7684\u6587\u4EF6\uFF0C\u540E\u7EED\u51E0\u79D2\u5B8C\u6210",
]

for i, feat in enumerate(features):
    y = Inches(2.5) + i * Inches(1.0)
    add_text_box(slide, Inches(1.2), y, Inches(10), Inches(0.8),
                 "\u2713  " + feat, font_size=24, color=TEXT_WHITE)

# ============================================================
# SLIDE 9: Core Feature 2 - Wiki Browser + Search
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_rounded_rect(slide, Inches(0.8), Inches(0.5), Inches(2.0), Inches(0.45),
                 fill_color=RGBColor(0xDD, 0xDB, 0xE5),
                 border_color=OBS_PURPLE_DIM,
                 text="\u6838\u5FC3\u529F\u80FD", font_size=13, text_color=OBS_PURPLE)

add_text_box(slide, Inches(0.8), Inches(1.2), Inches(11), Inches(0.8),
             "Wiki \u6D4F\u89C8\u5668 -- \u4F60\u7684\u77E5\u8BC6\u5361\u7247\u7D22\u5F15",
             font_size=36, color=TEXT_BRIGHT, bold=True)

add_decor_line(slide, Inches(0.8), Inches(2.0), Inches(2), OBS_PURPLE_DIM)

features = [
    "\u5361\u7247\u5F0F\u7D22\u5F15\uFF0C\u6309\u6982\u5FF5\u3001\u5B9E\u4F53\u3001\u7D20\u6750\u5206\u7C7B\u6D4F\u89C8",
    "\u70B9\u51FB\u5361\u7247\u67E5\u770B\u5B9A\u4E49\u3001\u8981\u70B9\u3001\u81EA\u52A8\u751F\u6210\u7684\u53CC\u5411\u94FE\u63A5",
    "\u6BCF\u4E2A\u9875\u9762\u81EA\u52A8\u5173\u8054\u76F8\u5173\u6982\u5FF5\uFF0C\u4E0D\u4F1A\u4EA7\u751F\u5B64\u5C9B",
    "\u53CD\u5411\u94FE\u63A5\uFF1A\u770B\u5230\u54EA\u4E9B\u9875\u9762\u5F15\u7528\u4E86\u8FD9\u4E2A\u6982\u5FF5",
]

for i, feat in enumerate(features):
    y = Inches(2.5) + i * Inches(1.0)
    add_text_box(slide, Inches(1.2), y, Inches(10), Inches(0.8),
                 "\u2713  " + feat, font_size=24, color=TEXT_WHITE)

# ============================================================
# SLIDE 10: Core Feature - Health Check (Lint)
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_rounded_rect(slide, Inches(0.8), Inches(0.5), Inches(2.0), Inches(0.45),
                 fill_color=RGBColor(0xDD, 0xDB, 0xE5),
                 border_color=OBS_PURPLE_DIM,
                 text="\u6838\u5FC3\u529F\u80FD", font_size=13, text_color=OBS_PURPLE)

add_text_box(slide, Inches(0.8), Inches(1.2), Inches(11), Inches(0.8),
             "\u77E5\u8BC6\u5065\u5EB7\u5EA6\u68C0\u67E5 -- \u7ED9\u4F60\u7684\u77E5\u8BC6\u5E93\u201C\u4F53\u68C0\u201D",
             font_size=36, color=TEXT_BRIGHT, bold=True)

add_text_box(slide, Inches(0.8), Inches(2.0), Inches(11), Inches(0.6),
             "\u5BF9\u5E94 Karpathy \u7684 lint \u64CD\u4F5C\uFF0C\u81EA\u52A8\u626B\u63CF\u77E5\u8BC6\u5E93\u53D1\u73B0\u95EE\u9898",
             font_size=18, color=TEXT_DIM)

features = [
    "\u68C0\u6D4B\u5B64\u5C9B\u9875\u9762\uFF1A\u6CA1\u6709\u53CC\u5411\u94FE\u63A5\u7684\u72EC\u7ACB\u6982\u5FF5",
    "\u68C0\u6D4B\u6B7B\u94FE\uFF1A\u94FE\u63A5\u6307\u5411\u4E0D\u5B58\u5728\u7684\u9875\u9762",
    "\u68C0\u6D4B\u6982\u5FF5\u7F3A\u53E3\uFF1A\u5DF2\u6709\u7D20\u6750\u4F46\u672A\u63D0\u70BC\u6210\u6982\u5FF5\u7684\u9886\u57DF",
    "\u4E00\u952E\u4FEE\u590D\u5EFA\u8BAE\uFF0C\u4FDD\u6301\u77E5\u8BC6\u5E93\u6301\u7EED\u5065\u5EB7",
]

for i, feat in enumerate(features):
    y = Inches(2.8) + i * Inches(1.0)
    add_text_box(slide, Inches(1.2), y, Inches(10), Inches(0.8),
                 "\u2713  " + feat, font_size=24, color=TEXT_WHITE)

# ============================================================
# SLIDE 11: Pro Feature - AI Chat (was 10)
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_rounded_rect(slide, Inches(0.8), Inches(0.5), Inches(2.5), Inches(0.45),
                 fill_color=OBS_PURPLE_DIM, border_color=OBS_PURPLE,
                 text="Full Version", font_size=13, text_color=TEXT_BRIGHT)

add_text_box(slide, Inches(0.8), Inches(1.2), Inches(11), Inches(0.8),
             "AI \u5BF9\u8BDD -- \u548C\u4F60\u7684\u77E5\u8BC6\u5E93\u804A\u5929",
             font_size=36, color=TEXT_BRIGHT, bold=True)

add_decor_line(slide, Inches(0.8), Inches(2.0), Inches(2), OBS_PURPLE)

features = [
    "\u9488\u5BF9\u4F60\u7684\u77E5\u8BC6\u5E93\u63D0\u95EE\uFF0CAI \u4ECE\u4F60\u7F16\u8BD1\u8FC7\u7684\u77E5\u8BC6\u4E2D\u627E\u7B54\u6848",
    "\u8BED\u4E49\u641C\u7D22\uFF1A\u4E0D\u662F\u5173\u952E\u8BCD\u5339\u914D\uFF0C\u800C\u662F\u7406\u89E3\u4F60\u7684\u610F\u601D",
    "\u6BCF\u4E2A\u56DE\u7B54\u81EA\u52A8\u6807\u6CE8 [[wiki-links]] \u5F15\u7528\u6765\u6E90",
    "\u76F8\u5F53\u4E8E\u4F60\u6709\u4E86\u4E00\u4E2A\u201C\u8BFB\u8FC7\u4F60\u6240\u6709\u7B14\u8BB0\u201D\u7684\u7814\u7A76\u52A9\u624F",
]

for i, feat in enumerate(features):
    y = Inches(2.5) + i * Inches(1.0)
    add_text_box(slide, Inches(1.2), y, Inches(10), Inches(0.8),
                 "\u2713  " + feat, font_size=24, color=TEXT_WHITE)

# ============================================================
# SLIDE 12: Pro Feature - Auto Compile (was 11)
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_rounded_rect(slide, Inches(0.8), Inches(0.5), Inches(2.5), Inches(0.45),
                 fill_color=OBS_PURPLE_DIM, border_color=OBS_PURPLE,
                 text="Full Version", font_size=13, text_color=TEXT_BRIGHT)

add_text_box(slide, Inches(0.8), Inches(1.2), Inches(11), Inches(0.8),
             "\u81EA\u52A8\u7F16\u8BD1 -- \u4E22\u8FDB\u53BB\u5C31\u4E0D\u7528\u7BA1\u4E86",
             font_size=36, color=TEXT_BRIGHT, bold=True)

add_decor_line(slide, Inches(0.8), Inches(2.0), Inches(2), OBS_PURPLE)

features = [
    "\u76D1\u542C raw/ \u76EE\u5F55\u53D8\u5316\uFF0C\u6709\u65B0\u7D20\u6750\u81EA\u52A8\u7F16\u8BD1\uFF0C\u4E0D\u7528\u624B\u52A8\u70B9",
    "\u65E9\u4E0A\u4E22\u4E00\u7BC7\u6587\u7AE0\u8FDB\u53BB\uFF0C\u559D\u676F\u5496\u5561\u56DE\u6765\uFF0C\u77E5\u8BC6\u5E93\u5DF2\u7ECF\u66F4\u65B0\u597D\u4E86",
    "\u589E\u91CF\u7F16\u8BD1 + \u81EA\u52A8\u89E6\u53D1\uFF0C\u8D8A\u7528\u8D8A\u5FEB",
]

for i, feat in enumerate(features):
    y = Inches(2.5) + i * Inches(1.0)
    add_text_box(slide, Inches(1.2), y, Inches(10), Inches(0.8),
                 "\u2713  " + feat, font_size=24, color=TEXT_WHITE)

# ============================================================
# SLIDE 13: Ecosystem plugins (was 12) - collecting materials
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_text_box(slide, Inches(0), Inches(0.5), W, Inches(0.8),
             "\u7D20\u6750\u600E\u4E48\u6765\uFF1F\u914D\u5957\u63D2\u4EF6\u63A8\u8350",
             font_size=36, color=TEXT_BRIGHT, bold=True,
             alignment=PP_ALIGN.CENTER)

add_decor_line(slide, Inches(6.0), Inches(1.3), Inches(1.3), OBS_PURPLE)

add_text_box(slide, Inches(1), Inches(1.7), Inches(11), Inches(0.6),
             "\u8FD9\u4E9B Obsidian \u63D2\u4EF6\u5E2E\u4F60\u628A\u7D20\u6750\u81EA\u52A8\u6536\u96C6\u5230 raw/ \u76EE\u5F55\uFF0C\u518D\u7528 Second Brain \u7F16\u8BD1",
             font_size=18, color=TEXT_DIM, alignment=PP_ALIGN.CENTER)

plugins = [
    {"name": "Obsidian Web Clipper",
     "desc": "\u6D4F\u89C8\u5668\u63D2\u4EF6\uFF0C\u4E00\u952E\u526A\u85CF\u7F51\u9875\u6587\u7AE0\u5230 raw/"},
    {"name": "\u5C0F\u7EA2\u4E66 Importer",
     "desc": "\u5BFC\u5165\u5C0F\u7EA2\u4E66\u6536\u85CF\u7B14\u8BB0\u5230 Obsidian"},
    {"name": "Podwise",
     "desc": "\u81EA\u52A8\u6574\u7406\u64AD\u5BA2\u5185\u5BB9\uFF0C\u751F\u6210\u7ED3\u6784\u5316\u7B14\u8BB0"},
]

card_w = Inches(3.2)
card_h = Inches(3.2)
start_x = Inches(1.8)
gap = Inches(0.6)

for i, plugin in enumerate(plugins):
    x = start_x + i * (card_w + gap)
    y = Inches(2.6)
    add_rounded_rect(slide, x, y, card_w, card_h,
                     fill_color=BG_CARD, border_color=BORDER_DIM)

    add_circle(slide, x + Inches(1.15), y + Inches(0.3), Inches(0.8),
               fill_color=OBS_PURPLE_DIM,
               text=["W", "\u5C0F", "P"][i],
               font_size=20, text_color=TEXT_BRIGHT)

    add_text_box(slide, x, y + Inches(1.3), card_w, Inches(0.5),
                 plugin["name"], font_size=20, color=TEXT_BRIGHT, bold=True,
                 alignment=PP_ALIGN.CENTER)

    add_text_box(slide, x + Inches(0.2), y + Inches(1.9), card_w - Inches(0.4), Inches(1.2),
                 plugin["desc"], font_size=16, color=TEXT_DIM,
                 alignment=PP_ALIGN.CENTER)

add_text_box(slide, Inches(0), Inches(6.2), W, Inches(0.6),
             "\u7D20\u6750\u6536\u96C6 + Second Brain \u7F16\u8BD1 = \u4ECE\u6536\u96C6\u5230\u6574\u7406\u5168\u81EA\u52A8",
             font_size=20, color=BRAND_GREEN, bold=True,
             alignment=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 14: Transition to demo (was 13)
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_text_box(slide, Inches(0), Inches(2.5), W, Inches(1.2),
             "\u63A5\u4E0B\u6765\uFF0C\u76F4\u63A5\u770B\u6548\u679C",
             font_size=48, color=TEXT_BRIGHT, bold=True,
             alignment=PP_ALIGN.CENTER)

add_decor_line(slide, Inches(6.0), Inches(3.7), Inches(1.3), OBS_PURPLE)

add_text_box(slide, Inches(0), Inches(4.2), W, Inches(0.8),
             "\u5B9E\u64CD\u6F14\u793A\uFF1A\u5B89\u88C5 -> \u914D\u7F6E -> \u7F16\u8BD1 -> \u6D4F\u89C8 Wiki",
             font_size=22, color=TEXT_DIM, alignment=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 15: Pricing (was 14)
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_text_box(slide, Inches(0), Inches(0.4), W, Inches(0.8),
             "\u8D39\u7528\u8BF4\u660E",
             font_size=36, color=TEXT_BRIGHT, bold=True,
             alignment=PP_ALIGN.CENTER)

add_decor_line(slide, Inches(6.0), Inches(1.2), Inches(1.3), OBS_PURPLE)

# Trial card
trial_x = Inches(1.5)
trial_w = Inches(4.5)
trial_h = Inches(5.3)
add_rounded_rect(slide, trial_x, Inches(1.5), trial_w, trial_h,
                 fill_color=BG_CARD, border_color=BORDER_DIM)

add_text_box(slide, trial_x, Inches(1.7), trial_w, Inches(0.5),
             "\u5185\u6D4B\u671F\u95F4", font_size=28, color=TEXT_DIM, bold=True,
             alignment=PP_ALIGN.CENTER)

add_text_box(slide, trial_x, Inches(2.3), trial_w, Inches(0.8),
             "\u514D\u8D39\u4F53\u9A8C\u5168\u90E8\u529F\u80FD", font_size=36, color=TEXT_BRIGHT, bold=True,
             alignment=PP_ALIGN.CENTER)

add_text_box(slide, trial_x, Inches(3.1), trial_w, Inches(0.5),
             "\u9996\u6B21\u7F16\u8BD1\u540E 14 \u5929\u5168\u529F\u80FD\u4F53\u9A8C", font_size=16, color=TEXT_MUTED,
             alignment=PP_ALIGN.CENTER)

trial_features = [
    "AI \u7F16\u8BD1\uFF08\u5168\u91CF / \u5355\u6587\u4EF6\uFF09",
    "Wiki \u6D4F\u89C8\u5668 + AI \u5BF9\u8BDD",
    "\u81EA\u52A8\u7F16\u8BD1 + \u8BED\u4E49\u641C\u7D22",
    "\u591A LLM \u652F\u6301 + \u591A\u8BED\u8A00",
    "\u77E5\u8BC6\u5065\u5EB7\u5EA6\u68C0\u67E5",
]
for i, feat in enumerate(trial_features):
    add_text_box(slide, trial_x + Inches(0.5), Inches(3.8) + i * Inches(0.52),
                 trial_w - Inches(1), Inches(0.5),
                 "\u2713  " + feat, font_size=15, color=TEXT_DIM)

# Buy card
buy_x = Inches(7.0)
buy_w = Inches(4.8)
buy_h = Inches(5.3)
add_rounded_rect(slide, buy_x, Inches(1.5), buy_w, buy_h,
                 fill_color=RGBColor(0xDD, 0xDB, 0xE5),
                 border_color=OBS_PURPLE)

add_rounded_rect(slide, buy_x + Inches(0.3), Inches(1.3),
                 Inches(1.8), Inches(0.4),
                 fill_color=OBS_PURPLE,
                 text="\u4E00\u6B21\u4E70\u65AD", font_size=12, text_color=TEXT_BRIGHT)

add_text_box(slide, buy_x, Inches(1.8), buy_w, Inches(0.5),
             "\u5168\u529F\u80FD\u7248", font_size=28, color=OBS_PURPLE_BRIGHT, bold=True,
             alignment=PP_ALIGN.CENTER)

add_text_box(slide, buy_x, Inches(2.4), buy_w, Inches(0.8),
             "\u00A5 30", font_size=52, color=TEXT_BRIGHT, bold=True,
             alignment=PP_ALIGN.CENTER)

buy_features = [
    "\u8BD5\u7528\u671F\u5168\u90E8\u529F\u80FD\uFF0C\u65E0\u9650\u5236",
    "\u81EA\u52A8\u7F16\u8BD1 -- \u65B0\u7D20\u6750\u81EA\u52A8\u5904\u7406",
    "AI \u5BF9\u8BDD -- \u56DE\u7B54\u5F15\u7528 [[wiki-links]]",
    "\u4E00\u952E\u5207\u6362 LLM \u540E\u7AEF",
    "\u77E5\u8BC6\u5065\u5EB7\u5EA6\u68C0\u67E5",
    "\u4E0D\u641E\u8BA2\u9605\uFF0C\u4E70\u4E00\u6B21\u7528\u4E00\u8F88\u5B50",
]
for i, feat in enumerate(buy_features):
    add_text_box(slide, buy_x + Inches(0.5), Inches(3.4) + i * Inches(0.52),
                 buy_w - Inches(1), Inches(0.5),
                 "\u2713  " + feat, font_size=15, color=TEXT_WHITE)

# ============================================================
# SLIDE 16: API Cost (was 15)
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_text_box(slide, Inches(0), Inches(0.8), W, Inches(0.8),
             "AI \u670D\u52A1\u8D39\u7528",
             font_size=36, color=TEXT_BRIGHT, bold=True,
             alignment=PP_ALIGN.CENTER)

add_decor_line(slide, Inches(6.0), Inches(1.6), Inches(1.3), OBS_PURPLE)

cost_items = [
    {"label": "\u63D2\u4EF6\u8BD5\u7528", "value": "\u9996\u6B21\u7F16\u8BD1\u540E 14 \u5929\u514D\u8D39"},
    {"label": "\u63D2\u4EF6\u4E70\u65AD", "value": "\u00A5 30 \u4E00\u6B21\u4E70\u65AD"},
    {"label": "DeepSeek API\uFF08\u63A8\u8350\uFF09", "value": "\u5145\u503C 10 \u5143\u7528\u51E0\u4E2A\u6708"},
    {"label": "\u6570\u636E\u5B89\u5168", "value": "\u5168\u90E8\u7559\u5728\u672C\u5730\uFF0C\u4E0D\u4E0A\u4F20"},
]

for i, item in enumerate(cost_items):
    y = Inches(2.2) + i * Inches(1.1)
    add_rounded_rect(slide, Inches(2.5), y, Inches(3.5), Inches(0.8),
                     fill_color=BG_CARD, border_color=BORDER_DIM,
                     text=item["label"], font_size=20, text_color=TEXT_DIM)
    add_text_box(slide, Inches(6.5), y + Inches(0.1), Inches(4.5), Inches(0.6),
                 item["value"], font_size=22, color=BRAND_GREEN, bold=True)

# ============================================================
# SLIDE 17: Call to action (was 16)
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_text_box(slide, Inches(0), Inches(0.8), W, Inches(0.8),
             "\u5F00\u59CB\u4F7F\u7528",
             font_size=36, color=TEXT_BRIGHT, bold=True,
             alignment=PP_ALIGN.CENTER)

add_decor_line(slide, Inches(6.0), Inches(1.6), Inches(1.3), OBS_PURPLE)

steps_cta = [
    {"num": "1", "text": "从 second-brain-release 的 Releases 下载 zip（GitHub 或 Gitee）"},
    {"num": "2", "text": "\u89E3\u538B\u540E\u7528 Obsidian \u6253\u5F00"},
    {"num": "3", "text": "\u914D\u7F6E DeepSeek API Key"},
    {"num": "4", "text": "\u5F00\u59CB\u7F16\u8BD1"},
]

for i, step in enumerate(steps_cta):
    y = Inches(2.2) + i * Inches(1.0)
    add_circle(slide, Inches(3.0), y, Inches(0.7),
               fill_color=OBS_PURPLE_DIM, text=step["num"],
               font_size=22, text_color=TEXT_BRIGHT)
    add_text_box(slide, Inches(4.0), y + Inches(0.1), Inches(6), Inches(0.6),
                 step["text"], font_size=24, color=TEXT_WHITE)

# ============================================================
# SLIDE 18: Links & Thanks (was 17)
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_text_box(slide, Inches(0), Inches(1.5), W, Inches(1.0),
             "Second Brain",
             font_size=48, color=TEXT_BRIGHT, bold=True,
             alignment=PP_ALIGN.CENTER)

add_decor_line(slide, Inches(6.0), Inches(2.5), Inches(1.3), OBS_PURPLE)

add_text_box(slide, Inches(0), Inches(3.0), W, Inches(0.6),
             "\u4E0B\u8F7D\u5730\u5740\u5728\u7B80\u4ECB\u548C\u8BC4\u8BBA\u533A\u7F6E\u9876",
             font_size=22, color=TEXT_DIM, alignment=PP_ALIGN.CENTER)

add_text_box(slide, Inches(0), Inches(3.8), W, Inches(0.6),
             "\u7231\u53D1\u7535: ifdian.net/a/ruowenwang",
             font_size=18, color=TEXT_MUTED, alignment=PP_ALIGN.CENTER)

add_text_box(slide, Inches(0), Inches(5.2), W, Inches(0.6),
             "\u89C9\u5F97\u6709\u7528\uFF0C\u8BF7\u4E00\u952E\u4E09\u8FDE\u652F\u6301\u4E00\u4E0B",
             font_size=24, color=OBS_PURPLE_BRIGHT, bold=True,
             alignment=PP_ALIGN.CENTER)

# ============================================================
# Save
# ============================================================
output_path = os.path.join(os.path.dirname(__file__), "SecondBrain-Bilibili.pptx")
prs.save(output_path)
print(f"PPT saved to: {output_path}")
